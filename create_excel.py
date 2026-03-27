import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Розничные расходы"

header_fill = PatternFill("solid", fgColor="4472C4")
white_font = Font(color="FFFFFF", bold=True)
thin = Side(style='thin')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

# Заголовок таблицы
ws.merge_cells("A1:F1")
ws["A1"] = "Розничные расходы в Европе"
ws["A1"].font = Font(bold=True, italic=True, size=13)
ws["A1"].alignment = Alignment(horizontal="center")

# Шапка таблицы (строки 2-3)
ws.merge_cells("A2:A3")
ws.merge_cells("B2:B3")
ws.merge_cells("C2:C3")
ws.merge_cells("D2:E2")
ws.merge_cells("F2:F3")

for ref, val in [
    ("A2", "Страна"),
    ("B2", "Население (тыс.)"),
    ("C2", "ВВП (млн)"),
    ("D2", "Общая сумма (млн)"),
    ("D3", "Продукты питания"),
    ("E3", "Одежда"),
    ("F2", "Соотношение\n(Питание/Одежда)"),
]:
    c = ws[ref]
    c.value = val
    c.font = white_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = border

# Данные
data = [
    ("Англия",     56753, 737886,  65506,  28262),
    ("Дания",       5134, 117238,  10582,   3246),
    ("Франция",    55385, 935579, 105354,  41515),
    ("Италия",     57221, 757177, 128082,  29050),
    ("Нидерланды", 14704, 253206,  21773,   7937),
]

for i, row in enumerate(data, start=4):
    for col, val in enumerate(row, start=1):
        c = ws.cell(row=i, column=col, value=val)
        c.border = border
        c.alignment = Alignment(horizontal="center")
    # Формула соотношения: Продукты питания / Одежда
    ratio = ws.cell(row=i, column=6)
    ratio.value = f"=D{i}/E{i}"
    ratio.number_format = "0.000"
    ratio.border = border
    ratio.alignment = Alignment(horizontal="center")

# Выделяем строку Италии (строка 7) зелёным — максимальное соотношение
for col in range(1, 7):
    ws.cell(row=7, column=col).fill = PatternFill("solid", fgColor="C6EFCE")
    ws.cell(row=7, column=col).font = Font(bold=True)

# Раздел с ответом
ws["A10"] = "Ответ:"
ws["A10"].font = Font(bold=True, size=12)

ws["A11"] = "Максимальное соотношение:"
ws["B11"] = "=MAX(F4:F8)"
ws["B11"].number_format = "0.000"
ws["B11"].font = Font(bold=True)

ws["A12"] = "Страна с макс. соотношением:"
ws["B12"] = "=INDEX(A4:A8,MATCH(MAX(F4:F8),F4:F8,0))"
ws["B12"].font = Font(bold=True, color="FF0000")

# Ширина столбцов и высота строк
for col, width in zip("ABCDEF", [16, 16, 14, 18, 14, 22]):
    ws.column_dimensions[col].width = width
ws.row_dimensions[2].height = 30
ws.row_dimensions[3].height = 30

wb.save("розничные_расходы.xlsx")
print("Файл создан: розничные_расходы.xlsx")
